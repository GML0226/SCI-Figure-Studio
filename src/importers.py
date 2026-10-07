"""Local source/data import. No current-workspace assumptions."""
from __future__ import annotations

import contextlib
import csv
import io
import os
import runpy
import sys
from pathlib import Path

import numpy as np
from matplotlib.figure import Figure
from matplotlib.backends.backend_agg import FigureCanvasAgg
import matplotlib.pyplot as plt

from core import Document

COLORS = ["#96C3F5", "#5B9FEF", "#BF83EC", "#9766CA", "#75BAAD", "#DE3F23", "#BDC3CD"]
HATCHES = ["", "/", "\\", "//", "..", "xx", ""]


def table_preview(path, sheet=None):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        names = book.sheetnames
        worksheet = book[sheet] if sheet else book.worksheets[0]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        book.close()
    else:
        names = []
        content = None
        for encoding in ("utf-8-sig", "gb18030", "utf-16"):
            try: content = path.read_text(encoding=encoding); break
            except UnicodeError: pass
        if content is None: raise ValueError("无法识别表格文件编码。")
        try: dialect = csv.Sniffer().sniff(content[:4096], delimiters=",\t;")
        except csv.Error: dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
        rows = list(csv.reader(io.StringIO(content), dialect))
    rows = [row for row in rows if any(value is not None and str(value).strip() for value in row)]
    if len(rows) < 2: raise ValueError("表格至少需要一行标题和一行数据。")
    return rows, names


def number(value):
    if value is None or str(value).strip().upper() in {"", "NULL", "NA", "N/A", "NAN", "NONE"}: return np.nan
    text = str(value).strip()
    return float(text[:-1]) / 100 if text.endswith("%") else float(text)


def create_chart(rows, category=0, columns=None, kind="bar", percent=False, series_rows=False, title="", ylabel="Value", size=(89, 65)):
    headers = [str(value) if value is not None else f"列 {i + 1}" for i, value in enumerate(rows[0])]
    if columns is None: columns = [i for i in range(len(headers)) if i != category]
    labels, values = [], []
    for row in rows[1:]:
        if category >= len(row) or row[category] is None or not str(row[category]).strip(): continue
        labels.append(str(row[category]))
        converted = []
        for column in columns:
            try: converted.append(number(row[column]) if column < len(row) else np.nan)
            except ValueError: converted.append(np.nan)
        values.append(converted)
    data = np.asarray(values, dtype=float)
    if not labels or not np.isfinite(data).any(): raise ValueError("所选数值列中没有可绘制的数值。")
    legends = [headers[i] for i in columns]
    if series_rows:
        data = data.T
        labels, legends = legends, labels
    if percent: data *= 100
    fig = Figure(figsize=(size[0] / 25.4, size[1] / 25.4), dpi=100)
    ax = fig.add_axes([0.16, 0.15, 0.80, 0.63])
    x = np.arange(len(labels), dtype=float)
    if kind in ("scatter", "line"):
        try: x = np.asarray([float(label) for label in labels])
        except ValueError: pass
    for index, legend in enumerate(legends):
        color, hatch = COLORS[index % len(COLORS)], HATCHES[index % len(HATCHES)]
        vals = data[:, index]
        valid = np.isfinite(vals)
        if kind in ("bar", "barh"):
            width = 0.8 / len(legends)
            positions = x + (index - (len(legends) - 1) / 2) * width
            if kind == "bar": ax.bar(positions[valid], vals[valid], width=width, label=legend, color=color, edgecolor="#303030", hatch=hatch, linewidth=.65)
            else: ax.barh(positions[valid], vals[valid], height=width, label=legend, color=color, edgecolor="#303030", hatch=hatch, linewidth=.65)
        elif kind == "scatter": ax.scatter(x[valid], vals[valid], label=legend, color=color, s=16)
        else: ax.plot(x, vals, label=legend, color=color, marker="o", ms=3, lw=1)
    if kind == "barh": ax.set_yticks(x, labels); ax.set_xlabel(ylabel)
    else: ax.set_xticks(x, labels); ax.set_ylabel(ylabel)
    ax.set_title(title, fontsize=9)
    ax.tick_params(labelsize=7)
    ax.grid(axis="x" if kind == "barh" else "y", color="#E4E4E4", lw=.4)
    ax.set_axisbelow(True)
    fig.legend(*ax.get_legend_handles_labels(), loc="upper center", bbox_to_anchor=(.53, .98), ncols=2, frameon=False, fontsize=7)
    FigureCanvasAgg(fig).draw()
    doc = Document([fig], [title or "导入表格"])
    doc.metadata["data"] = {"headers": headers, "rows": rows[1:], "category": category, "columns": columns}
    return doc


def load_python(path):
    """Run a plotting source locally, capturing figures instead of export/show."""
    path = Path(path).resolve()
    old_cwd, old_argv, old_path = Path.cwd(), sys.argv[:], sys.path[:]
    before = set(plt.get_fignums())
    captured = []
    show, close, savefig = plt.show, plt.close, Figure.savefig
    stdout = io.StringIO()
    import matplotlib
    old_rc = matplotlib.rcParams.copy()
    def capture(fig, *args, **kwargs):
        if fig not in captured: captured.append(fig)
    plt.show = lambda *args, **kwargs: [capture(plt.figure(number)) for number in plt.get_fignums() if number not in before]
    plt.close = lambda *args, **kwargs: None
    Figure.savefig = capture
    namespace = {}
    try:
        os.chdir(path.parent)
        sys.argv = [str(path)]
        sys.path.insert(0, str(path.parent))
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stdout):
            namespace = runpy.run_path(str(path), run_name="__main__")
            if not captured and callable(namespace.get("build_figure")):
                result = namespace["build_figure"]()
                candidates = [result] if isinstance(result, Figure) else list(result) if isinstance(result, (tuple, list)) else []
                for fig in candidates:
                    if isinstance(fig, Figure): capture(fig)
        for number in plt.get_fignums():
            if number not in before: capture(plt.figure(number))
        for value in namespace.values():
            if isinstance(value, Figure): capture(value)
        if not captured: raise ValueError("源文件没有创建 Matplotlib Figure。可在脚本中绘图，或提供 build_figure() 函数。")
    finally:
        plt.show, plt.close, Figure.savefig = show, close, savefig
        matplotlib.rcParams.update(old_rc)
        os.chdir(old_cwd); sys.argv = old_argv; sys.path[:] = old_path
    doc = Document(captured, [f"{path.stem} · 图 {i + 1}" for i in range(len(captured))])
    doc.metadata.update({"source": str(path), "import_log": stdout.getvalue()[-5000:]})
    return doc


def load_image(path):
    from PIL import Image
    raster = np.asarray(Image.open(path).convert("RGBA"))
    h, w = raster.shape[:2]
    fig = Figure(figsize=(89 / 25.4, 89 * h / w / 25.4))
    ax = fig.add_axes([0, 0, 1, 1]); ax.imshow(raster); ax.set_axis_off()
    doc = Document([fig], [Path(path).stem]); doc.metadata["source"] = str(path)
    return doc


def blank_document():
    fig = Figure(figsize=(89 / 25.4, 65 / 25.4))
    ax = fig.add_axes([.16, .16, .78, .75]); ax.set_xlabel("X"); ax.set_ylabel("Y")
    return Document([fig], ["新建科研图"])


def demo_document():
    fig = Figure(figsize=(160 / 25.4, 90 / 25.4))
    ax = fig.add_axes([.1, .18, .38, .64])
    x = np.arange(3)
    ax.bar(x - .17, [56, 68, 79], .34, label="Baseline", color=COLORS[0], edgecolor="#303030")
    ax.bar(x + .17, [72, 80, 87], .34, label="Our method", color=COLORS[5], edgecolor="#303030", hatch="xx")
    ax.set_xticks(x, ["Set A", "Set B", "Set C"]); ax.set_ylabel("Accuracy (%)")
    ax.set_title("Method comparison", fontsize=9)
    ax.legend(title="Methods", loc="upper left", fontsize=7, title_fontsize=8)
    bx = fig.add_axes([.6, .18, .36, .64])
    xx = np.arange(1, 6)
    bx.plot(xx, [1, 2, 3.5, 4, 5.5], "o-", color=COLORS[1], label="Series A")
    bx.scatter(xx, [2, 2.2, 3.8, 4.9, 5.7], color=COLORS[3], s=16, label="Series B")
    bx.set_xlabel("Step"); bx.set_ylabel("Response"); bx.set_title("Trend and observations", fontsize=9)
    bx.annotate("Editable note", (3, 3.5), xytext=(1.5, 5.6), fontsize=8, arrowprops={"arrowstyle": "->"})
    fig.text(.08, .92, "Scientific Figure Studio", fontsize=11, fontweight="bold")
    return Document([fig], ["功能示例"])
